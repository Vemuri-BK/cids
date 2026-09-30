"""Content of the 'prompt_guide' sheet: what every prompt column is and how it is made.

Examples are pulled from the real prompt table at build time, so they always
match the data.
"""
from __future__ import annotations

import pandas as pd

HEADER = ["Family", "Column", "What it contains", "How it is generated",
          "Patients", "Count", "Used for", "Example"]

TRAIN_EX = "DUKE_001"   # a training patient
TEST_EX = "DUKE_019"    # a test patient (has counterfactuals)


def _get(df: pd.DataFrame, pid: str, col: str) -> str:
    r = df.loc[df["patient_id"] == pid, col]
    v = "" if r.empty else str(r.iloc[0])
    return v if v and v.lower() != "nan" else "(filled after the Kaggle run)"


def guide_rows(df: pd.DataFrame) -> list[list[str]]:
    n_all = len(df)
    n_test = int((df["split"] == "test").sum())
    rows = [
        # ---------------- structured ----------------
        ["1. Structured (Type-S)", "prompt_S",
         "Main clinical instruction: Patient sentence (age bin, menopausal status, ethnicity) + "
         "Tumor sentence (subtype) + Acquisition sentence (vendor, field strength, view, laterality, "
         "phases, fat suppression, implant) + Task sentence.",
         "scripts/build_prompts.py: metadata cleaned (cids/prompts/fields.py), then a fixed template "
         "is filled (cids/prompts/structured.py). No AI; the same patient always gets the same text.",
         "all", str(n_all), "Main model; experiments E1, E3, E4, E5",
         _get(df, TRAIN_EX, "prompt_S")],
        ["1. Structured (Type-S)", "prompt_S_hints",
         "prompt_S plus rule-based radiology hints inserted before the task sentence.",
         "Same template + fixed rules: pre/peri-menopausal or young -> background-enhancement hint; "
         "post-menopausal -> low-background hint; subtype -> margin hint; 1.5T/3T -> noise hint; "
         "no fat suppression, implant, bilateral scan -> matching hints. Wording to be checked by a radiologist.",
         "all", str(n_all), "Ablation: does added radiology knowledge help?",
         _get(df, TRAIN_EX, "prompt_S_hints")],
        ["1. Structured (Type-S)", "prompt_S_nosubtype",
         "prompt_S without the tumor-subtype sentence.",
         "Same template with the Tumor sentence switched off.",
         "all", str(n_all), "Ablation: is subtype information needed, or a shortcut?",
         _get(df, TRAIN_EX, "prompt_S_nosubtype")],
        ["1. Structured (Type-S)", "prompt_S_noethnicity",
         "prompt_S without ethnicity.",
         "Same template with ethnicity switched off.",
         "all", str(n_all), "Fairness ablation",
         _get(df, TRAIN_EX, "prompt_S_noethnicity")],
        # ---------------- narrative ----------------
        ["2. Narrative (Type-N)", "narrative_1, narrative_2, narrative_3",
         "Free-text radiology-request style rewrites of prompt_S with exactly the same facts.",
         "scripts/generate_narratives.py on Kaggle GPU with Qwen2.5-7B-Instruct: 5 rewrites per patient, "
         "keep the best 3 that pass (1) fact check - every fact present, none invented/contradicted, no "
         "treatment/outcome words; (2) diversity - Jaccard distance to prompt_S in 0.30-0.50 (fallback 0.20-0.70); "
         "(3) same meaning - embedding cosine similarity >= 0.90. Ranked by 0.4*diversity + 0.6*similarity "
         "(FairVLM SRCP recipe + our fact check).",
         "all", f"{n_all} x 3", "Training/testing on narrative wording; experiments E2, E3, E4",
         _get(df, TRAIN_EX, "narrative_1")],
        ["2. Narrative (Type-N)", "narrative_heldout_1, _2, _3",
         "Same idea, but written by a DIFFERENT LLM.",
         "Same script and filters with Phi-3.5-mini-instruct. Never used in training.",
         "test only", f"{n_test} x 3", "Robustness to unseen wording (held-out prompt test)",
         _get(df, TEST_EX, "narrative_heldout_1")],
        # ---------------- counterfactual ----------------
        ["3. Counterfactual", "cf_age",
         "prompt_S with the age bin flipped (young <-> old) and menopausal status adjusted to match.",
         "scripts/build_prompts.py: copy the fields, change one, re-fill the template (seed 42).",
         "test only", str(n_test), "Demographic invariance: the mask should NOT change",
         _get(df, TEST_EX, "cf_age")],
        ["3. Counterfactual", "cf_ethnicity",
         "prompt_S with ethnicity swapped to a different group.",
         "Same as above; the new group is drawn at random (seed 42).",
         "test only", str(n_test), "Fairness test: the mask should NOT change",
         _get(df, TEST_EX, "cf_ethnicity")],
        ["3. Counterfactual", "cf_field_strength",
         "prompt_S with 1.5T <-> 3T swapped.",
         "Same as above.",
         "test only", str(n_test), "Acquisition sensitivity: the mask MAY change (shows the text is used)",
         _get(df, TEST_EX, "cf_field_strength")],
        # ---------------- dropout ----------------
        ["4. Field dropout", "(not stored)",
         "Any field replaced by 'unknown' with probability 0.3.",
         "Applied randomly by the training dataloader every batch, so it is not saved in this file.",
         "train", "random", "Robustness to missing metadata",
         "Patient aged 35-45, unknown menopausal status, African American. ..."],
    ]
    return rows


NOTES = [
    ["Reference patients", f"Training examples use {TRAIN_EX}; test-only examples (held-out, counterfactuals) use {TEST_EX}."],
    ["Excluded fields", "pCR, treatment, surgery and follow-up (outcome leakage); multifocal_cancer (partly derived "
                        "from the segmentation); breast_density (96% missing); oncotype_score and nottingham_grade "
                        "(87-99% missing). Full list in the field_dictionary sheet."],
    ["Missing values", "Missing metadata is written as 'unknown' (e.g. 'unknown menopausal status')."],
    ["Joining the image cache", "cache_id = lower-case patient_id, e.g. duke_001 -> train/duke_001.npz in the FADC 2-channel cache."],
    ["Splits", "split column = official MAMA-MIA train_test_splits.csv (1200 train / 306 test), identical to the FADC default split."],
    ["Filling narratives", "After the Kaggle run: python scripts/update_excel.py --csv cids_prompts_final.csv "
                           "--xlsx data/prompts/cids_prompts.xlsx"],
]


def write_guide_sheet(xw, df: pd.DataFrame) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    ws = xw.book.create_sheet("prompt_guide", 0)
    ws["A1"] = "CIDS prompt guide - what each prompt column is and how it was generated"
    ws["A1"].font = Font(bold=True, size=14, color="1F3864")
    ws.append([])
    ws.append(HEADER)
    hdr_row = ws.max_row
    fills = {"1.": "DDEBF7", "2.": "E2EFDA", "3.": "FCE4D6", "4.": "EDEDED"}
    thin = Side(style="thin", color="BFBFBF")
    for r in guide_rows(df):
        ws.append(r)
        fill = PatternFill("solid", fgColor=fills[r[0][:2]])
        for c in ws[ws.max_row]:
            c.fill = fill
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    for c in ws[hdr_row]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.append([])
    ws.append(["Notes"])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=12, color="1F3864")
    for k, v in NOTES:
        ws.append([k, v])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        ws.cell(row=ws.max_row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=ws.max_row, start_column=2, end_row=ws.max_row, end_column=8)
        ws.row_dimensions[ws.max_row].height = 32
    for col, w in zip("ABCDEFGH", (20, 24, 42, 58, 10, 10, 30, 70)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = ws.cell(row=hdr_row + 1, column=1)
