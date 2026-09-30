"""Build the CIDS prompt sheet (Type-S structured prompts + counterfactuals).

One row per patient (1506 rows). Narrative (Type-N) columns are created empty
and filled later by scripts/generate_narratives.py on Kaggle.

Usage (from the repo root):
    python scripts/build_prompts.py \
        --clinical "C:/Users/bhara/Desktop/MAMA_MIA_COMPLETE/clinical_and_imaging_info.xlsx" \
        --splits   "C:/Users/bhara/Desktop/MAMA_MIA_COMPLETE/train_test_splits.csv" \
        --out      data/prompts/cids_prompts.xlsx
"""
from __future__ import annotations

import argparse
import random
import sys
from pathlib import Path

import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from cids.prompts import (COUNTERFACTUAL_KINDS, EXCLUDED_FIELDS,  # noqa: E402
                          counterfactual_fields, normalize_row, structured_prompt)

N_NARRATIVES = 3
SEED = 42


def load_splits(path: str) -> dict:
    s = pd.read_csv(path)
    split = {}
    for pid in s["train_split"].dropna():
        split[str(pid).strip().upper()] = "train"
    for pid in s["test_split"].dropna():
        split[str(pid).strip().upper()] = "test"
    return split


def build(clinical: str, splits: str) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    df = pd.read_excel(clinical, sheet_name="dataset_info")
    split_of = load_splits(splits)
    rng = random.Random(SEED)
    rows = []
    for _, r in df.iterrows():
        raw = r.to_dict()
        pid = str(raw["patient_id"]).strip().upper()
        f = normalize_row(raw)
        split = split_of.get(pid, "unassigned")
        row = {
            "patient_id": pid,
            "cache_id": pid.lower(),            # name used in the FADC .npz cache
            "dataset": raw["dataset"],
            "split": split,
            "age": raw.get("age"),
            **f,
            "prompt_S": structured_prompt(f),
            "prompt_S_hints": structured_prompt(f, hints=True),
            "prompt_S_nosubtype": structured_prompt(f, with_subtype=False),
            "prompt_S_noethnicity": structured_prompt(f, with_ethnicity=False),
        }
        for i in range(1, N_NARRATIVES + 1):
            row[f"narrative_{i}"] = ""
        for i in range(1, N_NARRATIVES + 1):
            row[f"narrative_heldout_{i}"] = ""
        for kind in COUNTERFACTUAL_KINDS:
            if split == "test":
                row[f"cf_{kind}"] = structured_prompt(counterfactual_fields(f, kind, rng))
            else:
                row[f"cf_{kind}"] = ""
        rows.append(row)
    out = pd.DataFrame(rows)

    # ---- field dictionary ----
    used = [
        ("age", "age_bin", "Patient sentence", "demographic"),
        ("menopause", "menopause", "Patient sentence", "demographic"),
        ("ethnicity", "ethnicity", "Patient sentence (dropped in prompt_S_noethnicity)", "demographic"),
        ("tumor_subtype", "subtype", "Tumor sentence (dropped in prompt_S_nosubtype)", "clinical (from pre-treatment biopsy)"),
        ("manufacturer", "manufacturer", "Acquisition sentence", "acquisition"),
        ("field_strength", "field_strength", "Acquisition sentence", "acquisition"),
        ("view", "view", "Acquisition sentence", "acquisition"),
        ("bilateral_mri", "laterality", "Acquisition sentence", "acquisition"),
        ("fat_suppressed", "fat_suppressed", "Acquisition sentence", "acquisition"),
        ("has_implant", "implant", "Acquisition sentence", "acquisition"),
        ("num_phases", "num_phases", "Acquisition sentence", "acquisition"),
    ]
    fd = pd.DataFrame(
        [dict(source_column=a, prompt_field=b, used_in=c, role=d, status="USED") for a, b, c, d in used]
        + [dict(source_column=k, prompt_field="", used_in="", role="", status=f"EXCLUDED: {v}")
           for k, v in EXCLUDED_FIELDS.items()]
    )

    # ---- summary ----
    summ = []
    summ.append(("patients", len(out)))
    for s, n in out["split"].value_counts().items():
        summ.append((f"split={s}", n))
    for s, n in out["dataset"].value_counts().items():
        summ.append((f"dataset={s}", n))
    for col in ["age_bin", "menopause", "ethnicity_group", "subtype", "field_strength", "manufacturer"]:
        for v, n in out[col].value_counts().items():
            summ.append((f"{col}={v}", n))
    summ.append(("prompt_S words (mean)", round(out["prompt_S"].str.split().str.len().mean(), 1)))
    summ.append(("prompt_S_hints words (mean)", round(out["prompt_S_hints"].str.split().str.len().mean(), 1)))
    summ.append(("unique prompt_S strings", out["prompt_S"].nunique()))
    sm = pd.DataFrame(summ, columns=["item", "value"])
    return out, fd, sm


def write_excel(out: pd.DataFrame, fd: pd.DataFrame, sm: pd.DataFrame, path: Path) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        out.to_excel(xw, sheet_name="prompts", index=False)
        fd.to_excel(xw, sheet_name="field_dictionary", index=False)
        sm.to_excel(xw, sheet_name="summary", index=False)
        head_fill = PatternFill("solid", fgColor="1F3864")
        for name, frame in (("prompts", out), ("field_dictionary", fd), ("summary", sm)):
            ws = xw.sheets[name]
            ws.freeze_panes = "B2"
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = head_fill
                c.alignment = Alignment(wrap_text=True, vertical="center")
            for i, col in enumerate(frame.columns, start=1):
                letter = ws.cell(row=1, column=i).column_letter
                is_text = col.startswith(("prompt_", "narrative_", "cf_")) or col in ("used_in", "status")
                ws.column_dimensions[letter].width = 70 if is_text else max(12, min(28, len(str(col)) + 4))
                if is_text:
                    for cell in ws[letter][1:]:
                        cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.auto_filter.ref = ws.dimensions


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--clinical", required=True)
    ap.add_argument("--splits", required=True)
    ap.add_argument("--out", default="data/prompts/cids_prompts.xlsx")
    a = ap.parse_args()
    out, fd, sm = build(a.clinical, a.splits)
    p = Path(a.out)
    write_excel(out, fd, sm, p)
    out.to_csv(p.with_suffix(".csv"), index=False, encoding="utf-8")
    print(sm.to_string(index=False))
    print(f"\nWrote {p} and {p.with_suffix('.csv')}")


if __name__ == "__main__":
    main()
